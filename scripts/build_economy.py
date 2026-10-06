"""Official Sri Lanka economy series -> data/economy.json

Reads the Central Bank of Sri Lanka statistical-table spreadsheets (the file
names change with every update, so the links are found by their labels):
  * Interest Rates - Monthly (table 4.04): OPR, SDFR, SLFR, T-bill yields, AWPR, AWDR
  * Reserve Data Template - Latest (2.15.1): official reserve assets
  * Workers' Remittances (2.14.2) and Earnings from Tourism (2.14.1): monthly USD mn
  * Exports / Imports - Monthly (2.02 / 2.04): total exports and imports, USD mn
  * GDP by Industrial Origin at Constant 2015 Prices: quarterly real GDP -> growth
  * Household Population, Labour Force, Employment and Unemployment: unemployment rate
Each part is optional: if a sheet moves or changes shape, its previous values are kept.
Needs: pip install openpyxl
"""
import datetime as dt, html, io, json, os, re, subprocess

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'data', 'economy.json')
UA = 'Mozilla/5.0 (compatible; CeylonPulse-data/1.0; +https://github.com/investwithdon7-rgb/ceylonpulse-data)'
BASE = 'https://www.cbsl.gov.lk/en/statistics/statistical-tables/'
PAGES = {
    'monetary': BASE + 'monetary-sector',
    'external': BASE + 'external-sector',
    'accounts': BASE + 'real-sector/national-accounts',
    'labour': BASE + 'real-sector/prices-wages-employment',
}
MONTHS = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September', 'October', 'November', 'December']
MON3 = {m[:3]: i + 1 for i, m in enumerate(MONTHS)}
START_YEAR = 2015


def curl(url, binary=False):
    r = subprocess.run(['curl', '-sSL', '--fail', '-m', '120', '-A', UA, url], capture_output=True)
    if r.returncode != 0:
        print('  fetch failed', url[-80:])
        return None
    return r.stdout if binary else r.stdout.decode('utf-8', errors='replace')


_links = {}


def sheet_url(page, label_re):
    if page not in _links:
        raw = curl(PAGES[page]) or ''
        _links[page] = [(href, re.sub(r'\s+', ' ', html.unescape(re.sub(r'<[^>]+>', '', lab))).strip())
                        for href, lab in re.findall(r'(?is)<a[^>]+href="([^"]*sheets/[^"]+\.xlsx?)"[^>]*>(.*?)</a>', raw)]
    for href, lab in _links[page]:
        if re.search(label_re, lab, re.I):
            return href if href.startswith('http') else 'https://www.cbsl.gov.lk' + href
    print(f'  ! no sheet matching {label_re!r} on {page}')
    return None


def rows_of(url, sheet=None):
    data = curl(url, binary=True) if url else None
    if not data:
        return None
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def f(v):
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return None


def interest_rates():
    rows = rows_of(sheet_url('monetary', r'Interest Rates\s*-\s*Monthly'))
    if not rows:
        return None
    out, year = [], None
    for r in rows:
        if len(r) < 25:
            continue
        if isinstance(r[1], (int, float)) and 1990 < r[1] < 2100:
            year = int(r[1])
        m = r[2].strip() if isinstance(r[2], str) else None
        if year and m in MONTHS and year >= START_YEAR:
            # columns: 3 OPR, 4 SDFR, 5 SLFR, 7 T-bill 3M, 8 T-bill 12M, 15 AWDR, 24 AWPR
            out.append([f'{year}-{MONTHS.index(m) + 1:02d}', f(r[3]), f(r[4]), f(r[5]), f(r[7]), f(r[8]), f(r[24]), f(r[15])])
    if len(out) < 24:
        return None
    return {'columns': ['month', 'opr', 'sdfr', 'slfr', 'tbill3m', 'tbill12m', 'awpr', 'awdr'], 'months': out}


def year_by_month(label_re):
    """Sheets laid out as years across, months down (remittances, tourism earnings)."""
    rows = rows_of(sheet_url('external', label_re))
    if not rows:
        return None
    def year_of(x):
        # Header cells are numbers, or text such as "2026 (b)(c)" for provisional years
        m = re.match(r'\s*(20\d\d)\b', str(x)) if x is not None else None
        return int(m.group(1)) if m else None
    head = next((r for r in rows if sum(year_of(x) is not None for x in r) >= 5), None)
    if not head:
        return None
    cols = {j: year_of(x) for j, x in enumerate(head) if year_of(x) is not None}
    out = []
    for r in rows:
        name = next((x.strip() for x in r[:2] if isinstance(x, str)), '')
        if name in MONTHS:
            for j, y in cols.items():
                v = f(r[j]) if j < len(r) else None
                if v is not None and y >= START_YEAR:
                    out.append([f'{y}-{MONTHS.index(name) + 1:02d}', v])
    out.sort()
    return out or None


def reserves():
    rows = rows_of(sheet_url('external', r'Reserve Data Template\s*-\s*Latest'))
    if not rows:
        return None
    as_of = val = None
    for r in rows:
        for x in r:
            if isinstance(x, str):
                m = re.match(r'\s*End\s+([A-Za-z]+)\s+(\d{4})', x)
                if m and m.group(1) in MONTHS and not as_of:
                    as_of = f'{m.group(2)}-{MONTHS.index(m.group(1)) + 1:02d}'
        label = ' '.join(str(x) for x in r[:1] if x)
        if 'Official reserve assets' in label and val is None:
            for x in r[1:]:
                n = re.match(r'\s*([\d,]+(?:\.\d+)?)', str(x)) if x is not None else None
                if n:
                    val = float(n.group(1).replace(',', ''))
                    break
    return {'asOf': as_of, 'usdMn': val} if as_of and val else None


def trade_totals(page_label, sheet_hint, row_label):
    url = sheet_url('external', page_label)
    data = curl(url, binary=True) if url else None
    if not data:
        return None
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    name = next((s for s in wb.sheetnames if s.startswith(sheet_hint) and 'USD' in s and '2007' in s), None)
    if not name:
        return None
    rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
    head = next((r for r in rows if r and isinstance(r[0], str) and r[0].strip() == 'Category'), None)
    total = next((r for r in rows if r and isinstance(r[0], str) and r[0].strip().lower().startswith(row_label)), None)
    if not head or not total:
        return None
    out = {}
    for j, h in enumerate(head[1:], 1):
        ym = None
        if isinstance(h, dt.datetime):
            ym = f'{h.year}-{h.month:02d}'
        elif isinstance(h, str):
            m = re.match(r'([A-Z][a-z]{2})-(\d{2})', h.strip())
            if m and m.group(1) in MON3:
                ym = f'20{m.group(2)}-{MON3[m.group(1)]:02d}'
        v = f(total[j]) if j < len(total) else None
        if ym and v is not None and int(ym[:4]) >= START_YEAR:
            out[ym] = v
    return out or None


def gdp_growth():
    rows = rows_of(sheet_url('accounts', r'Industrial Origin.*Constant'))
    if not rows:
        return None
    head = next((r for r in rows if r and isinstance(r[0], str) and r[0].strip().startswith('Economic Activity')), None)
    gdp = next((r for r in rows if r and isinstance(r[0], str) and 'Gross Domestic Product (GDP)' in r[0] and 'Market Price' in r[0]), None)
    if not head or not gdp:
        return None
    q = {}
    for j, h in enumerate(head):
        m = re.match(r'\s*(\d{4})\s*Quarter\s*(\d)', str(h or ''), re.S)
        if m and j < len(gdp) and f(gdp[j]):
            q[(int(m.group(1)), int(m.group(2)))] = float(gdp[j])
    out = [[f'{y} Q{k}', round((v / q[(y - 1, k)] - 1) * 100, 1)] for (y, k), v in sorted(q.items()) if (y - 1, k) in q]
    return out or None


def unemployment():
    rows = rows_of(sheet_url('labour', r'Household Population.*Unemployment'))
    if not rows:
        return None
    head = next((r for r in rows if r and any(isinstance(x, str) and 'Unemployment Rate' in x for x in r)), None)
    if not head:
        return None
    col = next(j for j, x in enumerate(head) if isinstance(x, str) and 'Unemployment Rate' in x)
    out, section = [], None
    for r in rows:
        first = str(r[0]).strip() if r and r[0] is not None else ''
        filled = [str(x).strip() for x in r if x is not None and str(x).strip()]
        if len(filled) == 1 and not re.match(r'\d{4}', filled[0]):
            section = filled[0]   # "All", then other breakdowns (label can sit in any column)
            continue
        if section == 'All' and re.match(r'\d{4}( Q\d)?$', first) and col < len(r) and f(r[col]) is not None:
            if ' Q' in first and int(first[:4]) >= START_YEAR:
                out.append([first, round(float(r[col]), 1)])
    return out or None


def main():
    old = json.load(open(OUT, encoding='utf-8')) if os.path.exists(OUT) else {}
    old.pop('generated', None)
    payload = dict(old)
    parts = {
        'rates': interest_rates,
        'reserves': reserves,
        'remittances': lambda: year_by_month(r"Workers.? Remittances \(2009"),
        'tourismEarnings': lambda: year_by_month(r'Earnings from Tourism'),
        'gdpGrowth': gdp_growth,
        'unemployment': unemployment,
    }
    for key, fn in parts.items():
        try:
            v = fn()
        except Exception as e:  # one broken sheet must not stop the others
            print(f'  {key}: error {e}')
            v = None
        if v:
            payload[key] = v
            print(f'  {key}: ok ({len(v) if isinstance(v, list) else "latest"})')
        else:
            print(f'  {key}: kept previous')
    try:
        ex = trade_totals(r'Exports\s*-\s*Monthly', '2.02', 'total exports')
        im = trade_totals(r'Imports\s*-\s*Monthly', '2.04', 'total imports')
        if ex and im:
            payload['trade'] = [[m, ex[m], im.get(m)] for m in sorted(ex)]
            print(f'  trade: ok ({len(payload["trade"])})')
    except Exception as e:
        print(f'  trade: error {e}')
    payload['source'] = 'Central Bank of Sri Lanka statistical tables; Department of Census and Statistics (via CBSL)'
    payload['url'] = BASE
    if payload == old:
        print('economy.json: unchanged')
        return
    out = {'generated': dt.datetime.now(dt.timezone.utc).isoformat(timespec='seconds'), **payload}
    json.dump(out, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
    print(f'economy.json: updated ({os.path.getsize(OUT)} bytes)')


if __name__ == '__main__':
    main()
